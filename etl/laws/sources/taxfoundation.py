"""
Tax Foundation — income tax, sales tax, and grocery treatment.

Tax Foundation publishes each table as an .xlsx linked from a data page. We
parse the spreadsheet (exact numbers, no scraping of prose) and cite the page.

    income_tax_top_rate, income_tax_structure
        "State Individual Income Tax Rates and Brackets" (as of Jan 1) and
        Facts & Figures Table 11 (as of late April). Both are read; for each
        state the LATER-dated table is used, so a spring rate change wins.
    sales_tax_combined
        "State and Local Sales Tax Rates" — the midyear edition (July 1) when
        it exists, otherwise the January one.
    grocery_tax_exempt
        Facts & Figures Table 31, "Sales Tax Treatment of Groceries".

Editions move to new URLs each year, so each is looked up by year, newest
first. A page that exists but has no spreadsheet, or a spreadsheet whose
layout changed, raises SourceError and the law keeps its old (aging) rows.
"""

from __future__ import annotations

import io
import re
from datetime import date

import openpyxl

from ..common import Fact, SourceError, SourceResult, STATES, clean, get, parse_long_date, state_code

NAME = "Tax Foundation"
BASE = "https://taxfoundation.org/data/all/state"


def _xlsx_links(page: str) -> list[str]:
    return list(dict.fromkeys(re.findall(r'https://[^"\'\s>]+\.xlsx', page)))


def _first_edition(candidates: list[str]) -> tuple[str, str]:
    """(page_url, page_html) for the first candidate page that loads and links a spreadsheet."""
    errors = []
    for url in candidates:
        try:
            page = get(url)
        except SourceError as exc:
            errors.append(str(exc))
            continue
        if _xlsx_links(page):
            return url, page
        errors.append(f"{url}: no .xlsx link")
    raise SourceError("no edition found: " + "; ".join(errors))


def _workbook(url: str) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(get(url, binary=True)), data_only=True)


def _as_of(text: str, fallback: date) -> date:
    m = re.search(r"as of\s+([A-Z][a-z]+ \d{1,2}, 20\d\d)", text, re.I)
    return (parse_long_date(m.group(1)) if m else None) or fallback


def _pct(x: float) -> float:
    return round(x * 100, 3)


def _fmt_pct(x: float) -> str:
    return f"{_pct(x):g}%"


# ---------------------------------------------------------------------------
# Income tax
# ---------------------------------------------------------------------------


def parse_income_blocks(rows: list[tuple]) -> dict[str, dict]:
    """State -> {"rates": [single-filer marginal rates], "label", "none", "capital_gains_only"}.

    Layout (both the standalone table and F&F Table 11): a state's label in
    column A starts a block; each bracket row has `rate, '>', threshold` with
    the single-filer rate first. 'none' means no wage income tax. Footnote rows
    at the bottom have text in column A only and are ignored.
    """
    # Footnotes: "(k) In Wash., tax rates apply only to high earners' capital
    # gains income." A state whose label cites such a footnote has no wage
    # income tax, whatever its rate rows say.
    cap_gains_notes = set()
    for row in rows:
        text = row[0] if row else None
        m = re.match(r"\(([a-z]{1,3})\)\s*(.*)", text.strip()) if isinstance(text, str) else None
        if m and re.search(r"capital gains income", m.group(2), re.I) and re.search(r"\bonly\b", m.group(2), re.I):
            cap_gains_notes.add(m.group(1))

    blocks: dict[str, dict] = {}
    cur: dict | None = None
    for row in rows:
        row = tuple(row)
        label = row[0] if row else None
        code = state_code(label) if isinstance(label, str) and len(label) < 40 else None
        if code:
            cur = blocks.setdefault(code, {"rates": [], "label": label.strip(),
                                           "none": False, "capital_gains_only": False})
            cited = set(re.findall(r"\b([a-z]{1,3})\b", " ".join(re.findall(r"\((.*?)\)", label))))
            if cited & cap_gains_notes:
                cur["capital_gains_only"] = True
        if cur is None:
            continue
        if isinstance(label, str) and label.strip().startswith("(") and len(label) > 40:
            cur = None  # footnotes start: stop attaching rows to the last state
            continue
        cells = list(row[1:])
        texts = " ".join(str(c) for c in cells if isinstance(c, str)).lower()
        if "capital gains" in texts:
            cur["capital_gains_only"] = True
        if re.search(r"\bnone\b", texts) and not any(isinstance(c, (int, float)) for c in cells[:1]):
            cur["none"] = True
        for i in range(len(cells) - 1):
            if isinstance(cells[i], (int, float)) and cells[i + 1] == ">":
                cur["rates"].append(float(cells[i]))
                break  # first rate in the row is the single filer's
    return blocks


def _income_facts(blocks: dict[str, dict], *, page: str, data_url: str, as_of: date,
                  table: str) -> dict[str, tuple[Fact, Fact]]:
    out = {}
    for code, b in blocks.items():
        if b["none"] or b["capital_gains_only"] or not b["rates"]:
            top, structure = 0.0, "none"
        else:
            top = max(b["rates"])
            structure = "flat" if len(set(b["rates"])) == 1 else "graduated"
        rates = ", ".join(_fmt_pct(r) for r in b["rates"]) or "none"
        note = ""
        if b["capital_gains_only"]:
            note = f"No tax on wage income; tax on capital gains income only ({rates})."
        quote = f"{table}: {b['label']} — single filer rates: {rates}"
        rate_fact = Fact(value=f"{_pct(top):g}", value_numeric=_pct(top), source_name=NAME,
                         source_url=page, data_url=data_url, source_date=as_of,
                         quote=quote, notes=note)
        struct_fact = Fact(value=structure, source_name=NAME, source_url=page,
                           data_url=data_url, source_date=as_of, quote=quote, notes=note)
        out[code] = (rate_fact, struct_fact)
    return out


def income_tax(today: date) -> tuple[SourceResult, SourceResult]:
    candidates = []
    # Standalone table (as of Jan 1).
    for year in (today.year, today.year - 1):
        candidates.append(("standalone", f"{BASE}/state-income-tax-rates-{year}/", year))
    # Facts & Figures Table 11 (as of late April).
    for year in (today.year, today.year - 1):
        candidates.append(("ff", f"{BASE}/{year}-state-tax-data/", year))

    tables: list[tuple[date, dict]] = []
    notes = []
    seen_kinds = set()
    for kind, url, year in candidates:
        if kind in seen_kinds:
            continue
        try:
            page = get(url)
        except SourceError as exc:
            notes.append(str(exc))
            continue
        links = _xlsx_links(page)
        if kind == "standalone":
            links = [x for x in links if "Income-Tax" in x or "income-tax" in x.lower()]
        if not links:
            notes.append(f"{url}: no spreadsheet")
            continue
        wb = _workbook(links[0])
        if kind == "standalone":
            ws = wb.worksheets[0]  # newest year is the first sheet
            as_of = _as_of(clean(page), date(year, 1, 1))
            label = f"Tax Foundation {ws.title} state income tax rates table"
        else:
            if "11" not in wb.sheetnames:
                notes.append(f"{links[0]}: no Table 11 sheet")
                continue
            ws = wb["11"]
            head = " ".join(str(c) for r in ws.iter_rows(max_row=4, values_only=True) for c in r if c)
            if "Income Tax Rates" not in head:
                raise SourceError(f"Facts & Figures sheet 11 is no longer the income tax table: {head[:80]}")
            as_of = _as_of(head, date(year, 1, 1))
            label = f"Facts & Figures {year}, Table 11"
        blocks = parse_income_blocks(list(ws.iter_rows(values_only=True)))
        if set(blocks) != STATES:
            raise SourceError(f"{label}: expected 51 jurisdictions, got {len(blocks)} "
                              f"(missing {sorted(STATES - set(blocks))})")
        tables.append((as_of, _income_facts(blocks, page=url, data_url=links[0],
                                            as_of=as_of, table=label)))
        seen_kinds.add(kind)

    if not tables:
        raise SourceError("no Tax Foundation income tax table found: " + "; ".join(notes))

    tables.sort(key=lambda t: t[0], reverse=True)   # newest first
    rate = SourceResult("income_tax_top_rate", notes=notes)
    struct = SourceResult("income_tax_structure", notes=list(notes))
    newest_date, newest = tables[0]
    for code in STATES:
        r, s = newest[code]
        # Where the older table disagrees, say so — it usually means a mid-year change.
        for older_date, older in tables[1:]:
            o = older[code][0]
            if o.value != r.value:
                r.notes = (r.notes + " " if r.notes else "") + (
                    f"Tax Foundation's {older_date:%B %-d, %Y} table listed {o.value}%.")
        rate.facts[code] = r
        struct.facts[code] = s
    return rate, struct


# ---------------------------------------------------------------------------
# Sales tax
# ---------------------------------------------------------------------------


def parse_sales_sheet(rows: list[tuple]) -> tuple[dict[str, tuple], str]:
    """State -> (state_rate, avg_local, combined); plus the title row."""
    rows = [tuple(r) for r in rows]
    title = " ".join(str(c) for r in rows[:2] for c in r if isinstance(c, str))
    header_i = next((i for i, r in enumerate(rows)
                     if r and any(isinstance(c, str) and "Combined" in c for c in r)), None)
    if header_i is None:
        raise SourceError("sales tax sheet has no 'Combined' column")
    header = [str(c or "") for c in rows[header_i]]
    col = {name: next(i for i, h in enumerate(header) if name in h)
           for name in ("State Tax Rate", "Average Local", "Combined Tax Rate")}
    out = {}
    for r in rows[header_i + 1:]:
        code = state_code(r[0]) if r and isinstance(r[0], str) else None
        if not code:
            continue
        vals = [r[col[k]] for k in ("State Tax Rate", "Average Local", "Combined Tax Rate")]
        out[code] = tuple(float(v) if isinstance(v, (int, float)) else 0.0 for v in vals)
    return out, title


def sales_tax(today: date) -> SourceResult:
    candidates = []
    for year in (today.year, today.year - 1):
        candidates += [f"{BASE}/{year}-sales-tax-rates-midyear/", f"{BASE}/sales-tax-rates-{year}/"]
    candidates.append(f"{BASE}/sales-tax-rates/")

    best = None
    notes = []
    for url in candidates:
        try:
            page = get(url)
        except SourceError as exc:
            notes.append(str(exc))
            continue
        links = [x for x in _xlsx_links(page) if "sales" in x.lower()]
        if not links:
            continue
        rows, title = parse_sales_sheet(list(_workbook(links[0]).worksheets[0].iter_rows(values_only=True)))
        as_of = _as_of(title, date(today.year, 1, 1))
        if best is None or as_of > best[0]:
            best = (as_of, url, links[0], rows, title)
        if best and best[0].year == today.year and best[0].month >= 7:
            break  # this year's midyear edition is the newest there is
    if best is None:
        raise SourceError("no Tax Foundation sales tax table found: " + "; ".join(notes))

    as_of, url, data_url, rows, title = best
    if set(rows) != STATES:
        raise SourceError(f"sales tax table: expected 51 jurisdictions, got {len(rows)}")
    res = SourceResult("sales_tax_combined", notes=notes)
    for code, (st, local, combined) in rows.items():
        quote = (f"{title}: {code} state rate {_fmt_pct(st)}, average local rate "
                 f"{_fmt_pct(local)}, combined {_fmt_pct(combined)}")
        res.facts[code] = Fact(value=f"{_pct(combined):g}", value_numeric=_pct(combined),
                               source_name=NAME, source_url=url, data_url=data_url,
                               source_date=as_of, quote=quote,
                               notes=f"State {_fmt_pct(st)} + average local {_fmt_pct(local)}.")
    return res


# ---------------------------------------------------------------------------
# Groceries (Facts & Figures Table 31)
# ---------------------------------------------------------------------------


def classify_grocery(state_rate: object, treatment: object) -> tuple[str, str]:
    """(value, note) per the rubric: exempt | reduced | taxed."""
    t = str(treatment).strip().lower() if treatment is not None else ""
    if t in ("--", "—", "n.a.", "none", ""):
        return "exempt", "No state sales tax."
    if t == "exempt":
        return "exempt", ""
    if "included in base" in t or "full" in t:
        return "taxed", "Groceries are taxed at the full state rate."
    if isinstance(treatment, (int, float)) and isinstance(state_rate, (int, float)):
        if treatment >= state_rate:
            return "taxed", f"Groceries taxed at {_fmt_pct(treatment)} (the full state rate)."
        return "reduced", f"Groceries taxed at {_fmt_pct(treatment)}, vs. {_fmt_pct(state_rate)} general rate."
    raise SourceError(f"unrecognized grocery treatment {treatment!r}")


def groceries(today: date) -> SourceResult:
    url, page = _first_edition([f"{BASE}/{y}-state-tax-data/" for y in (today.year, today.year - 1)])
    data_url = _xlsx_links(page)[0]
    wb = _workbook(data_url)
    if "31" not in wb.sheetnames:
        raise SourceError("Facts & Figures has no Table 31 sheet")
    rows = [tuple(r) for r in wb["31"].iter_rows(values_only=True)]
    head = " ".join(str(c) for r in rows[:4] for c in r if c)
    if "Groceries" not in head:
        raise SourceError(f"Facts & Figures sheet 31 is no longer the grocery table: {head[:80]}")
    as_of = _as_of(head, date(today.year, 1, 1))
    header_i = next(i for i, r in enumerate(rows) if r and r[0] == "State")
    res = SourceResult("grocery_tax_exempt")
    for r in rows[header_i + 1:]:
        code = state_code(r[0]) if isinstance(r[0], str) and len(r[0]) < 30 else None
        if not code:
            continue
        value, note = classify_grocery(r[1], r[2])
        shown = [str(r[0]).strip(), _fmt_pct(r[1]) if isinstance(r[1], float) else str(r[1]),
                 _fmt_pct(r[2]) if isinstance(r[2], float) else str(r[2])]
        res.facts[code] = Fact(value=value, source_name=NAME, source_url=url, data_url=data_url,
                               source_date=as_of,
                               quote=f"Facts & Figures Table 31 (as of {as_of:%B %-d, %Y}): "
                                     f"{shown[0]} — state general sales tax {shown[1]}, "
                                     f"grocery treatment {shown[2]}",
                               notes=note)
    if set(res.facts) != STATES:
        raise SourceError(f"grocery table: expected 51 jurisdictions, got {len(res.facts)}")
    return res
